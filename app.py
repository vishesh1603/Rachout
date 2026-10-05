import os
import re
import io
import json
import base64
import smtplib
import mimetypes
import zipfile
import xml.etree.ElementTree as ET
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from flask import Flask, request, jsonify, render_template, send_file

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

app = Flask(__name__)

# DuckDuckGo search
try:
    from duckduckgo_search import DDGS
except ImportError:
    DDGS = None

# New google.genai SDK
try:
    from google import genai
    from google.genai import types
except ImportError:
    genai = None
    types = None

# Document parsing: pypdf & python-docx
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

# ReportLab imports
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
except ImportError:
    pass

# Safe HRFlowable import
try:
    from reportlab.platypus import HRFlowable
except ImportError:
    try:
        from reportlab.platypus.flowables import HRFlowable
    except ImportError:
        from reportlab.platypus import Flowable
        class HRFlowable(Flowable):
            def __init__(self, width="100%", thickness=1, color=HexColor('#cccccc'), spaceBefore=1, spaceAfter=1):
                super().__init__()
                self.width = width
                self.thickness = thickness
                self.color = color
                self.spaceBefore = spaceBefore
                self.spaceAfter = spaceAfter
            def wrap(self, availWidth, availHeight):
                self.availWidth = availWidth
                return availWidth, self.thickness + self.spaceBefore + self.spaceAfter
            def draw(self):
                self.canv.saveState()
                self.canv.setStrokeColor(self.color)
                self.canv.setLineWidth(self.thickness)
                self.canv.line(0, self.spaceAfter, self.availWidth, self.spaceAfter)
                self.canv.restoreState()

def extract_file_content(file_bytes, filename, content_type=None):
    """
    Extracts text or prepares multimodal payload from various file types:
    PDF, Word (.docx, .doc), Images (.jpg, .jpeg, .png, .webp, .bmp), Text (.txt, .md, .rtf, .html, etc.)
    """
    ext = os.path.splitext(filename.lower())[1]
    text = ""
    is_image = False
    mime_type = content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"

    # 1. Images (JPG, PNG, WEBP, BMP, GIF)
    if ext in ['.jpg', '.jpeg', '.png', '.webp', '.bmp', '.gif'] or (content_type and content_type.startswith('image/')):
        is_image = True
        if ext in ['.jpg', '.jpeg']:
            mime_type = 'image/jpeg'
        elif ext == '.png':
            mime_type = 'image/png'
        elif ext == '.webp':
            mime_type = 'image/webp'
        elif ext == '.gif':
            mime_type = 'image/gif'
        text = ""

    # 2. PDF Documents
    elif ext == '.pdf' or mime_type == 'application/pdf':
        mime_type = 'application/pdf'
        if pypdf is not None:
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                pages_text = []
                for p in reader.pages:
                    extracted = p.extract_text()
                    if extracted:
                        pages_text.append(extracted)
                text = "\n".join(pages_text).strip()
            except Exception as e:
                print(f"Error extracting PDF text: {e}")

    # 3. Word Documents (.docx)
    elif ext == '.docx':
        mime_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        if docx is not None:
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                lines = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if row_line:
                            lines.append(row_line)
                text = "\n".join(lines).strip()
            except Exception as e:
                print(f"python-docx parsing error: {e}")
        if not text:
            # Fallback to direct XML extraction from docx zip
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                    xml_data = z.read('word/document.xml')
                    tree = ET.fromstring(xml_data)
                    text = "".join(tree.itertext()).strip()
            except Exception as ze:
                print(f"zipfile docx error: {ze}")

    # 4. Word Legacy Documents (.doc)
    elif ext == '.doc':
        mime_type = 'application/msword'
        matches = re.findall(rb'[\x20-\x7E\t\r\n]{4,}', file_bytes)
        text = "\n".join(m.decode('latin-1', errors='ignore') for m in matches if len(m.strip()) > 3).strip()

    # 5. Text / Markdown / HTML / RTF / JSON / CSV
    else:
        for enc in ['utf-8', 'latin-1', 'cp1252']:
            try:
                text = file_bytes.decode(enc).strip()
                if ext == '.rtf':
                    text = re.sub(r'\\[a-z0-9]+ ?', '', text)
                    text = re.sub(r'[{}\\]', '', text).strip()
                break
            except Exception:
                continue

    return {
        "text": text,
        "is_image": is_image,
        "mime_type": mime_type,
        "filename": filename,
        "size_bytes": len(file_bytes)
    }

def parse_gemini_json(text):
    text_clean = text.strip()
    
    # Try parsing directly
    try:
        return json.loads(text_clean)
    except json.JSONDecodeError:
        pass

    # Try matching ```json ... ``` blocks
    match = re.search(r'```(?:json)?\s*(\{.*\})\s*```', text_clean, re.DOTALL | re.IGNORECASE)
    if match:
        try:
            return json.loads(match.group(1))
        except json.JSONDecodeError:
            pass

    # Search for first '{' and last '}'
    first_brace = text_clean.find('{')
    last_brace = text_clean.rfind('}')
    if first_brace != -1 and last_brace != -1:
        try:
            return json.loads(text_clean[first_brace:last_brace+1])
        except json.JSONDecodeError:
            pass

    raise ValueError("Could not parse JSON from Gemini response")

def build_resume_pdf(resume_data, candidate_name):
    pdf_buffer = io.BytesIO()
    margin = 46.8  # 0.65 inch = 46.8 points
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    name_style = ParagraphStyle(
        'ResumeName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=4,
        textColor=HexColor('#111111')
    )
    
    contact_style = ParagraphStyle(
        'ResumeContact',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        alignment=TA_CENTER,
        spaceAfter=8,
        textColor=HexColor('#444444')
    )
    
    section_heading_style = ParagraphStyle(
        'ResumeSectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        spaceBefore=8,
        spaceAfter=2,
        textColor=HexColor('#222222')
    )
    
    summary_style = ParagraphStyle(
        'ResumeSummary',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        alignment=TA_JUSTIFY,
        spaceAfter=6,
        textColor=HexColor('#333333')
    )
    
    item_title_style = ParagraphStyle(
        'ResumeItemTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        textColor=HexColor('#222222')
    )
    
    item_date_style = ParagraphStyle(
        'ResumeItemDate',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=12,
        alignment=TA_RIGHT,
        textColor=HexColor('#555555')
    )
    
    bullet_style = ParagraphStyle(
        'ResumeBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=2,
        textColor=HexColor('#333333')
    )
    
    # 1. Name & Contact
    name = resume_data.get('name', candidate_name)
    story.append(Paragraph(str(name), name_style))
    
    contact = resume_data.get('contact', '')
    if contact:
        story.append(Paragraph(str(contact), contact_style))
    
    story.append(HRFlowable(width="100%", thickness=1, color=HexColor('#111111'), spaceBefore=2, spaceAfter=8))
    
    # Printable width: 518.4
    col_widths = [410, 108]
    
    def add_section_header(title):
        story.append(Paragraph(title.upper(), section_heading_style))
        story.append(HRFlowable(width="100%", thickness=0.5, color=HexColor('#888888'), spaceBefore=1, spaceAfter=6))
        
    # 2. Summary
    summary = resume_data.get('summary', '')
    if summary:
        add_section_header('Professional Summary')
        story.append(Paragraph(str(summary), summary_style))
        
    # 3. Experience
    experience = resume_data.get('experience', [])
    if experience:
        add_section_header('Work Experience')
        for exp in experience:
            if not isinstance(exp, dict):
                continue
            exp_story = []
            title = str(exp.get('title', ''))
            company = str(exp.get('company', ''))
            date = str(exp.get('date', ''))
            bullets = exp.get('bullets', [])
            if isinstance(bullets, str):
                bullets = [bullets]
            
            left_text = f"<b>{title}</b> &mdash; {company}" if company else f"<b>{title}</b>"
            t = Table([[Paragraph(left_text, item_title_style), Paragraph(date, item_date_style)]], colWidths=col_widths)
            t.setStyle(TableStyle([
                ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ('TOPPADDING', (0,0), (-1,-1), 2),
            ]))
            exp_story.append(t)
            
            for bullet in bullets:
                exp_story.append(Paragraph(f"&bull; {str(bullet)}", bullet_style))
                
            story.append(KeepTogether(exp_story))
            story.append(Spacer(1, 4))
            
    # 4. Projects
    projects = resume_data.get('projects', [])
    if projects:
        add_section_header('Key Projects')
        for proj in projects:
            if not isinstance(proj, dict):
                continue
            proj_story = []
            pname = str(proj.get('name', ''))
            tech = str(proj.get('tech', ''))
            bullets = proj.get('bullets', [])
            if isinstance(bullets, str):
                bullets = [bullets]
            
            left_text = f"<b>{pname}</b> &mdash; <i>{tech}</i>" if tech else f"<b>{pname}</b>"
            t = Table([[Paragraph(left_text, item_title_style), Paragraph("", item_date_style)]], colWidths=col_widths)
            t.setStyle(TableStyle([
                ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ('TOPPADDING', (0,0), (-1,-1), 2),
            ]))
            proj_story.append(t)
            
            for bullet in bullets:
                exp_story.append(Paragraph(f"&bull; {str(bullet)}", bullet_style))
                
            story.append(KeepTogether(proj_story))
            story.append(Spacer(1, 4))
            
    # 5. Skills
    skills = resume_data.get('skills', [])
    if skills:
        add_section_header('Skills')
        skills_story = []
        for skill in skills:
            if isinstance(skill, dict):
                cat = str(skill.get('category', ''))
                items = skill.get('items', '')
                if isinstance(items, list):
                    items = ', '.join(str(i) for i in items)
                skills_story.append(Paragraph(f"<b>{cat}:</b> {str(items)}", bullet_style))
            elif isinstance(skill, str):
                skills_story.append(Paragraph(f"&bull; {skill}", bullet_style))
        story.append(KeepTogether(skills_story))
        story.append(Spacer(1, 4))
        
    # 6. Education
    education = resume_data.get('education', [])
    if education:
        add_section_header('Education')
        for edu in education:
            if not isinstance(edu, dict):
                continue
            edu_story = []
            degree = str(edu.get('degree', ''))
            school = str(edu.get('school', ''))
            date = str(edu.get('date', ''))
            
            left_text = f"<b>{degree}</b> &mdash; {school}" if school else f"<b>{degree}</b>"
            t = Table([[Paragraph(left_text, item_title_style), Paragraph(date, item_date_style)]], colWidths=col_widths)
            t.setStyle(TableStyle([
                ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ('TOPPADDING', (0,0), (-1,-1), 2),
            ]))
            edu_story.append(t)
            story.append(KeepTogether(edu_story))
            story.append(Spacer(1, 4))
            
    doc.build(story)
    pdf_bytes = pdf_buffer.getvalue()
    pdf_buffer.close()
    return pdf_bytes

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/parse-file', methods=['POST'])
def parse_file():
    try:
        file_bytes = None
        filename = "uploaded_file"
        content_type = None

        if 'file' in request.files:
            uploaded = request.files['file']
            filename = uploaded.filename or "uploaded_file"
            file_bytes = uploaded.read()
            content_type = uploaded.content_type
        else:
            data = request.json or {}
            b64_data = data.get('data') or data.get('file_base64', '')
            filename = data.get('filename', 'uploaded_file')
            content_type = data.get('mime_type')
            if b64_data:
                if ',' in b64_data:
                    b64_data = b64_data.split(',', 1)[1]
                file_bytes = base64.b64decode(b64_data)

        if not file_bytes:
            return jsonify({"error": "No file content received."}), 400

        result = extract_file_content(file_bytes, filename, content_type)
        result["file_base64"] = base64.b64encode(file_bytes).decode('utf-8')
        return jsonify({"success": True, **result})
    except Exception as e:
        print(f"API parse-file error: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/analyze-ats', methods=['POST'])
def analyze_ats():
    try:
        data = request.json or {}
        resume_text = data.get('resume', '').strip()
        jd_text = data.get('jd', '').strip()
        resume_file = data.get('resume_file')
        jd_file = data.get('jd_file')
        role = data.get('role', 'Target Role').strip()
        company_name = data.get('company_name', 'Target Company').strip()
        candidate_name = data.get('candidate_name', 'Candidate').strip()

        # Process uploaded files for Resume and JD
        resume_parts = []
        if resume_file and isinstance(resume_file, dict) and resume_file.get('data'):
            r_b64 = resume_file.get('data', '')
            if ',' in r_b64:
                r_b64 = r_b64.split(',', 1)[1]
            r_bytes = base64.b64decode(r_b64)
            r_info = extract_file_content(r_bytes, resume_file.get('filename', 'resume'), resume_file.get('mime_type'))
            if r_info['is_image']:
                resume_parts.append(types.Part.from_bytes(data=r_bytes, mime_type=r_info['mime_type']))
            elif r_info['mime_type'] == 'application/pdf':
                resume_parts.append(types.Part.from_bytes(data=r_bytes, mime_type='application/pdf'))
                if r_info['text'] and not resume_text:
                    resume_text = r_info['text']
            elif r_info['text'] and not resume_text:
                resume_text = r_info['text']

        jd_parts = []
        if jd_file and isinstance(jd_file, dict) and jd_file.get('data'):
            j_b64 = jd_file.get('data', '')
            if ',' in j_b64:
                j_b64 = j_b64.split(',', 1)[1]
            j_bytes = base64.b64decode(j_b64)
            j_info = extract_file_content(j_bytes, jd_file.get('filename', 'job_description'), jd_file.get('mime_type'))
            if j_info['is_image']:
                jd_parts.append(types.Part.from_bytes(data=j_bytes, mime_type=j_info['mime_type']))
            elif j_info['mime_type'] == 'application/pdf':
                jd_parts.append(types.Part.from_bytes(data=j_bytes, mime_type='application/pdf'))
                if j_info['text'] and not jd_text:
                    jd_text = j_info['text']
            elif j_info['text'] and not jd_text:
                jd_text = j_info['text']

        if not resume_text and not resume_parts:
            return jsonify({"error": "Please provide your resume (paste text or upload a file)."}), 400
        if not jd_text and not jd_parts:
            return jsonify({"error": "Please provide a target job description (paste text or upload a file)."}), 400

        gemini_api_key = os.environ.get("GEMINI_API_KEY")
        if not gemini_api_key:
            return jsonify({"error": "GEMINI_API_KEY environment variable is not set. Please set it in your .env file or environment."}), 500

        if genai is None or types is None:
            return jsonify({"error": "google.genai library is not available."}), 500

        client = genai.Client(api_key=gemini_api_key)

        prompt_instructions = f"""You are a certified ATS (Applicant Tracking System) audit engine and senior talent acquisition director.

Candidate Name: {candidate_name}
Target Role: {role}
Target Company: {company_name}

Audit the candidate's CURRENT un-tailored resume strictly against the target Job Description.
Calculate a realistic, data-driven ATS match score (0-100) reflecting how an enterprise ATS (like Workday, Greenhouse, Lever) would rank this candidate right now.

Respond ONLY with valid JSON in this exact structure:
{{
  "current_score": 62,
  "score_breakdown": {{
    "skills_match": 58,
    "experience_match": 65,
    "formatting_score": 78
  }},
  "summary": "2-3 sentence candid executive diagnosis of the candidate's current alignment against this JD.",
  "matched_skills": ["Skill1", "Skill2", "Skill3", "Skill4", "Skill5"],
  "missing_skills": ["MissingCriticalSkill1", "MissingSkill2", "MissingSkill3", "MissingSkill4"],
  "strengths": [
    "Specific strong qualification or experience",
    "Relevant background element"
  ],
  "improvements": [
    "Specific action to incorporate missing keyword/skill",
    "Specific metric or quantification recommendation",
    "Structural or section prioritization recommendation"
  ]
}}
"""
        contents = [prompt_instructions]
        contents.append("\n=== CANDIDATE CURRENT RESUME ===\n")
        if resume_text:
            contents.append(f"Resume Text:\n{resume_text}\n")
        for p in resume_parts:
            contents.append(p)

        contents.append(f"\n=== TARGET JOB DESCRIPTION ({role} at {company_name}) ===\n")
        if jd_text:
            contents.append(f"Job Description Text:\n{jd_text}\n")
        for p in jd_parts:
            contents.append(p)

        config = types.GenerateContentConfig(response_mime_type="application/json")
        candidate_models = [
            os.environ.get("GEMINI_MODEL", "gemini-2.5-flash"),
            "gemini-2.0-flash",
            "gemini-1.5-flash"
        ]
        models_to_try = []
        for m in candidate_models:
            if m and m not in models_to_try:
                models_to_try.append(m)

        response_text = None
        last_err = None
        for model_name in models_to_try:
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=config
                )
                response_text = response.text
                if response_text:
                    break
            except Exception as e:
                print(f"ATS audit error on {model_name}: {e}")
                last_err = e

        if not response_text:
            return jsonify({"error": f"Gemini API error during ATS analysis: {str(last_err)}"}), 500

        result_data = parse_gemini_json(response_text)
        return jsonify(result_data)
    except Exception as e:
        print(f"API analyze-ats error: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/generate', methods=['POST'])
def generate():
    try:
        data = request.json or {}
        resume_text = data.get('resume', '').strip()
        jd_text = data.get('jd', '').strip()
        resume_file = data.get('resume_file')
        jd_file = data.get('jd_file')
        company_name = data.get('company_name', '').strip()
        recruiter_email = data.get('recruiter_email', '').strip()
        role = data.get('role', '').strip()
        candidate_name = data.get('candidate_name', '').strip()
        
        # 1. Company research via DuckDuckGo search
        company_context = ""
        if company_name and DDGS is not None:
            try:
                with DDGS() as ddgs:
                    results = list(ddgs.text(f"{company_name} company mission culture tech stack", max_results=5))
                company_context = "\n".join([r.get('body', '') for r in results if 'body' in r])
            except Exception as search_err:
                print(f"DuckDuckGo search error: {search_err}")
                company_context = ""

        # 2. Process uploaded files for Resume and JD
        resume_parts = []
        if resume_file and isinstance(resume_file, dict) and resume_file.get('data'):
            r_b64 = resume_file.get('data', '')
            if ',' in r_b64:
                r_b64 = r_b64.split(',', 1)[1]
            r_bytes = base64.b64decode(r_b64)
            r_info = extract_file_content(r_bytes, resume_file.get('filename', 'resume'), resume_file.get('mime_type'))
            if r_info['is_image']:
                resume_parts.append(types.Part.from_bytes(data=r_bytes, mime_type=r_info['mime_type']))
            elif r_info['mime_type'] == 'application/pdf':
                resume_parts.append(types.Part.from_bytes(data=r_bytes, mime_type='application/pdf'))
                if r_info['text'] and not resume_text:
                    resume_text = r_info['text']
            elif r_info['text'] and not resume_text:
                resume_text = r_info['text']

        jd_parts = []
        if jd_file and isinstance(jd_file, dict) and jd_file.get('data'):
            j_b64 = jd_file.get('data', '')
            if ',' in j_b64:
                j_b64 = j_b64.split(',', 1)[1]
            j_bytes = base64.b64decode(j_b64)
            j_info = extract_file_content(j_bytes, jd_file.get('filename', 'job_description'), jd_file.get('mime_type'))
            if j_info['is_image']:
                jd_parts.append(types.Part.from_bytes(data=j_bytes, mime_type=j_info['mime_type']))
            elif j_info['mime_type'] == 'application/pdf':
                jd_parts.append(types.Part.from_bytes(data=j_bytes, mime_type='application/pdf'))
                if j_info['text'] and not jd_text:
                    jd_text = j_info['text']
            elif j_info['text'] and not jd_text:
                jd_text = j_info['text']

        # Ensure we have at least text or parts for both
        if not resume_text and not resume_parts:
            return jsonify({"error": "Please provide a resume (paste text or upload PDF, Word, Image, or Text file)."}), 400
        if not jd_text and not jd_parts:
            return jsonify({"error": "Please provide a job description (paste text or upload PDF, Word, Image, or Text file)."}), 400

        # 3. Google GenAI Client
        gemini_api_key = os.environ.get("GEMINI_API_KEY")
        if not gemini_api_key:
            return jsonify({"error": "GEMINI_API_KEY environment variable is not set. Please add it to your .env file or environment."}), 500
            
        if genai is None or types is None:
            return jsonify({"error": "google.genai Python library is not installed or failed to import."}), 500
            
        client = genai.Client(api_key=gemini_api_key)
        
        prompt_instructions = f"""You are an elite career coach and ATS resume expert.

Company research context (from web search):
{company_context}

Candidate Name: {candidate_name}
Target Role: {role}
Target Company: {company_name}

Tasks:
1. Create an ATS-optimized 1-page resume tailored to the target Job Description. Mirror exact keywords and skills from the JD.
   Output structure must match this JSON schema:
   {{
     "name": "{candidate_name}",
     "contact": "email | phone | linkedin | location",
     "summary": "2-3 sentence impactful summary tailored to the role",
     "experience": [{{"title": "...", "company": "...", "date": "...", "bullets": ["..."]}}],
     "projects": [{{"name": "...", "tech": "...", "bullets": ["..."]}}],
     "skills": [{{"category": "...", "items": "..."}}],
     "education": [{{"degree": "...", "school": "...", "date": "..."}}]
   }}

2. Write a personalized cold outreach email to a recruiter at {company_name} for the role "{role}".
   - Reference 1-2 specific things from the company research (real, specific, not generic)
   - Under 200 words
   - Professional but direct tone, no filler phrases like "I hope this finds you well"
   - Start with "Subject: ..." on the first line
   - Sign off with "{candidate_name}"

Respond ONLY with valid JSON:
{{
  "optimized_ats_score": 96,
  "improvements_applied": [
    "Integrated missing keywords directly from the JD",
    "Quantified key accomplishments with impactful metrics",
    "Refined summary and skills to match role requirements"
  ],
  "resume": {{ ... }},
  "email": "Subject: ...\\n\\nHi [Recruiter's Name],\\n\\n...",
  "email_subject": "..."
}}
"""

        contents = [prompt_instructions]
        
        # Add Resume Content
        contents.append("\n=== CANDIDATE RESUME INPUT ===\n")
        if resume_text:
            contents.append(f"Resume Text:\n{resume_text}\n")
        for p in resume_parts:
            contents.append(p)

        # Add Job Description Content
        contents.append(f"\n=== TARGET JOB DESCRIPTION ({role} at {company_name}) ===\n")
        if jd_text:
            contents.append(f"Job Description Text:\n{jd_text}\n")
        for p in jd_parts:
            contents.append(p)

        config = types.GenerateContentConfig(
            response_mime_type="application/json"
        )

        candidate_models = [
            os.environ.get("GEMINI_MODEL", "gemini-2.5-flash"),
            "gemini-2.0-flash",
            "gemini-1.5-flash"
        ]
        models_to_try = []
        for m in candidate_models:
            if m and m not in models_to_try:
                models_to_try.append(m)

        response_text = None
        last_err = None

        for model_name in models_to_try:
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=config
                )
                response_text = response.text
                if response_text:
                    break
            except Exception as api_err:
                print(f"Gemini model ({model_name}) error: {api_err}. Trying fallback if available...")
                last_err = api_err

        if not response_text:
            return jsonify({"error": f"Gemini API error: {str(last_err)}"}), 500
        
        try:
            parsed_data = parse_gemini_json(response_text)
            return jsonify(parsed_data)
        except Exception as parse_err:
            print(f"Error parsing Gemini response: {parse_err}")
            return jsonify({
                "error": "Failed to parse tailored response as structured JSON from LLM.",
                "raw_response": response_text
            }), 500

    except Exception as e:
        print(f"API generate error: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/preview-pdf', methods=['POST'])
def preview_pdf():
    try:
        data = request.json or {}
        resume_data = data.get('resume', {})
        candidate_name = data.get('candidate_name', 'Candidate')
        
        if not resume_data:
            return jsonify({"error": "No resume data provided for PDF generation."}), 400
            
        pdf_bytes = build_resume_pdf(resume_data, candidate_name)
        pdf_b64 = base64.b64encode(pdf_bytes).decode('utf-8')
        
        return jsonify({"pdf_base64": pdf_b64})
    except Exception as e:
        print(f"API preview-pdf error: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/download-pdf', methods=['POST'])
def download_pdf():
    try:
        data = request.json or {}
        resume_data = data.get('resume', {})
        candidate_name = data.get('candidate_name', 'Candidate')
        role = data.get('role', 'Resume')
        
        if not resume_data:
            return jsonify({"error": "No resume data provided for download."}), 400
            
        pdf_bytes = build_resume_pdf(resume_data, candidate_name)
        safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', candidate_name)
        safe_role = re.sub(r'[^a-zA-Z0-9_]', '_', role)
        filename = f"{safe_name}_{safe_role}_ATS_Resume.pdf"
        
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype='application/pdf',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        print(f"API download-pdf error: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/send', methods=['POST'])
def send_email():
    try:
        data = request.json or {}
        gmail_user = data.get('gmail_user', '')
        gmail_app_password = data.get('gmail_app_password', '')
        to_email = data.get('to_email', '')
        subject = data.get('subject', '')
        email_body = data.get('email_body', '')
        resume_data = data.get('resume_data', {})
        candidate_name = data.get('candidate_name', 'Candidate')
        role = data.get('role', 'Candidate')
        
        if not gmail_user or not gmail_app_password or not to_email or not subject or not email_body or not resume_data:
            return jsonify({"error": "Missing required fields for sending email."}), 400
            
        # Re-generate PDF bytes
        pdf_bytes = build_resume_pdf(resume_data, candidate_name)
        
        # Build safe filename
        safe_name = re.sub(r'[^a-zA-Z0-9_]', '_', candidate_name)
        safe_role = re.sub(r'[^a-zA-Z0-9_]', '_', role)
        filename = f"{safe_name}_{safe_role}.pdf"
        
        # Compose MIME email
        msg = MIMEMultipart()
        msg['From'] = gmail_user
        msg['To'] = to_email
        msg['Subject'] = subject
        msg.attach(MIMEText(email_body, 'plain'))
        
        # Attach PDF
        part = MIMEBase('application', 'octet-stream')
        part.set_payload(pdf_bytes)
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', f'attachment; filename="{filename}"')
        msg.attach(part)
        
        # Send via Gmail SMTP_SSL
        with smtplib.SMTP_SSL('smtp.gmail.com', 465) as server:
            server.login(gmail_user, gmail_app_password)
            server.sendmail(gmail_user, to_email, msg.as_string())
            
        return jsonify({"success": True, "filename": filename})
    except Exception as e:
        print(f"API send error: {e}")
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
